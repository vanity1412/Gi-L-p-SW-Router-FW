from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = r"D:\temp\Lap03\Bao_cao_Lab03_Syslog_Zabbix_Chinh_xac_FINAL.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9E2EC", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def apply_table_style(table, widths=None, header_fill="E8EEF5"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for row_idx, row in enumerate(table.rows):
        if row_idx == 0:
            tr_pr = row._tr.get_or_add_trPr()
            tbl_header = tr_pr.find(qn("w:tblHeader"))
            if tbl_header is None:
                tbl_header = OxmlElement("w:tblHeader")
                tr_pr.append(tbl_header)
            tbl_header.set(qn("w:val"), "true")
        for cell_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if row_idx == 0:
                set_cell_shading(cell, header_fill)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True
            if widths and cell_idx < len(widths):
                cell.width = widths[cell_idx]


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    if level == 1:
        p.paragraph_format.keep_with_next = True
    return p


def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.style = "Normal"
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(text, style="List Bullet")
    return p


def add_number(doc, text):
    p = doc.add_paragraph(text, style="List Number")
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.right_indent = Cm(0.2)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(9)
    return p


def add_callout(doc, title, text, fill="F4F6F9"):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    set_table_borders(table, color="B8C6D6")
    p = cell.paragraphs[0]
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor(31, 58, 95)
    p.add_run(f" {text}")
    doc.add_paragraph()


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr[i].text = header
    for row_data in rows:
        row = table.add_row().cells
        for i, value in enumerate(row_data):
            row[i].text = str(value)
    apply_table_style(table, widths=widths)
    doc.add_paragraph()
    return table


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for style_name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Lab 03 - Centralized Syslog Monitoring with Zabbix")
    return doc


def build():
    doc = setup_document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("BÁO CÁO LAB 03")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(11, 37, 69)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("XÂY DỰNG HỆ THỐNG THU THẬP SYSLOG TẬP TRUNG\nVÀ CẢNH BÁO AN TOÀN MẠNG BẰNG ZABBIX")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(46, 116, 181)

    info_rows = [
        ("Môn học", "Hệ thống giám sát an toàn mạng"),
        ("Sinh viên thực hiện", "................................................"),
        ("Mã sinh viên", "................................................"),
        ("Lớp", "................................................"),
        ("Giảng viên hướng dẫn", "................................................"),
        ("Năm học", "2025 - 2026"),
    ]
    add_table(doc, ["Thông tin", "Nội dung"], info_rows, widths=[Inches(2.0), Inches(4.2)])

    add_callout(
        doc,
        "Tóm tắt:",
        "Bài lab xây dựng một hệ thống thu thập Syslog tập trung về Zabbix Server, "
        "áp dụng bộ lọc để phát hiện sự kiện bất thường và gửi cảnh báo qua Web, Email, "
        "Telegram/SMS gateway, audio frontend hoặc dịch vụ call như Twilio/PagerDuty.",
        fill="EAF4FF",
    )

    doc.add_page_break()

    add_heading(doc, "1. Giới thiệu bài lab", 1)
    add_body(
        doc,
        "Trong hệ thống mạng, server, switch, router và firewall liên tục sinh ra log về trạng thái hoạt động, "
        "lỗi hệ thống, đăng nhập, thay đổi cấu hình, trạng thái interface và các sự kiện an toàn thông tin. "
        "Nếu quản trị viên kiểm tra từng thiết bị thủ công, việc phát hiện sự cố sẽ chậm và dễ bỏ sót."
    )
    add_body(
        doc,
        "Bài lab này triển khai mô hình thu thập Syslog tập trung bằng Zabbix. Zabbix Server kết hợp rsyslog "
        "đóng vai trò máy nhận log trung tâm. Các thiết bị gửi Syslog về server qua UDP/TCP port 514. "
        "Zabbix đọc log, lọc sự kiện quan trọng, tạo Problem trên dashboard và kích hoạt Action gửi cảnh báo."
    )

    add_heading(doc, "2. Mục tiêu và yêu cầu", 1)
    for item in [
        "Cấu hình rsyslog trên Zabbix Server để nhận Syslog tập trung từ Linux Server, Switch, Router/Firewall.",
        "Lưu log theo từng hostname trong thư mục /var/log/remote/ để dễ truy vết.",
        "Tạo item kiểu log[] hoặc logrt[] trong Zabbix để đọc file log liên tục.",
        "Xây dựng rule/trigger để phát hiện interface down, login thất bại, lỗi hệ thống, cấu hình thay đổi và sự kiện bảo mật.",
        "Hiển thị cảnh báo trên Zabbix Problems/Dashboard và gửi thông báo qua Email, Telegram/SMS gateway, audio frontend hoặc call gateway.",
        "Kết hợp giám sát SNMP cho các chỉ số CPU, RAM, disk, sessions, traffic counters và load average khi cần demo metric thiết bị.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "3. Mô hình triển khai", 1)
    add_body(
        doc,
        "Mô hình logic của bài lab gồm ba lớp: nguồn log, máy thu thập trung tâm và lớp cảnh báo. "
        "Nguồn log có thể là Server Linux, Switch, Router hoặc Firewall. Zabbix Server nhận log qua rsyslog, "
        "sau đó Zabbix Agent đọc file log và Zabbix Server xử lý trigger."
    )
    add_code(
        doc,
        "Server / Switch / Router / Firewall\n"
        "        |\n"
        "        |-- Syslog UDP/TCP 514 --> rsyslog trên Zabbix Server\n"
        "        |                         /var/log/remote/<hostname>.log\n"
        "        |\n"
        "        |-- SNMP v2c 1161/161 --> Zabbix SNMP items (metric bổ sung)\n"
        "                                  CPU/RAM/Disk/Sessions/Traffic/Load\n"
        "        ↓\n"
        "Zabbix items -> Trigger rules -> Problems -> Actions -> Web/Email/Telegram/SMS/Call"
    )
    add_callout(
        doc,
        "Lưu ý kỹ thuật:",
        "Syslog và SNMP là hai luồng dữ liệu khác nhau. Syslog dùng để thu thập sự kiện/log. "
        "SNMP dùng để thu thập chỉ số định lượng. Trong lab demo, có thể dùng cả hai để bài trình bày đầy đủ hơn.",
        fill="FFF7E6",
    )

    add_heading(doc, "4. Bảng thiết bị và vai trò", 1)
    add_table(
        doc,
        ["Thiết bị", "Vai trò", "Dữ liệu gửi về Zabbix", "Ghi chú cấu hình"],
        [
            ("Zabbix Server", "Máy giám sát trung tâm", "Nhận Syslog, đọc log, tạo trigger, gửi alert", "Cài zabbix-server, zabbix-agent/agent2, rsyslog"),
            ("Linux Server", "Máy chủ dịch vụ", "auth.log, syslog, failed login, sudo, service error", "Gửi bằng rsyslog hoặc logger test"),
            ("Switch", "Thiết bị chuyển mạch", "Interface up/down, cấu hình thay đổi, warning/error", "logging host <IP_ZABBIX>"),
            ("Router/Firewall", "Gateway hoặc firewall giả lập", "Interface, NAT/gateway, deny, login fail, warning/error", "logging host <IP_ZABBIX>"),
            ("SNMPSim Dashboard", "Nguồn giả lập metric SNMP", "CPU, RAM, disk, sessions, traffic counters, load average", "Dùng port 1161 nếu snmpd chiếm port 161"),
        ],
        widths=[Inches(1.45), Inches(1.55), Inches(2.35), Inches(1.3)],
    )

    add_heading(doc, "5. Cấu hình Zabbix Server nhận Syslog tập trung", 1)
    add_heading(doc, "5.1 Cài đặt và bật rsyslog", 2)
    add_code(doc, "sudo apt update\nsudo apt install rsyslog -y\nsudo systemctl enable --now rsyslog")

    add_heading(doc, "5.2 Tạo cấu hình nhận log từ xa", 2)
    add_code(
        doc,
        "sudo nano /etc/rsyslog.d/10-remote-syslog.conf\n\n"
        "module(load=\"imudp\")\n"
        "input(type=\"imudp\" port=\"514\")\n\n"
        "module(load=\"imtcp\")\n"
        "input(type=\"imtcp\" port=\"514\")\n\n"
        "$template RemoteSyslog,\"/var/log/remote/%HOSTNAME%.log\"\n"
        "*.* ?RemoteSyslog\n"
        "& stop"
    )

    add_heading(doc, "5.3 Tạo thư mục lưu log và khởi động lại dịch vụ", 2)
    add_code(
        doc,
        "sudo mkdir -p /var/log/remote\n"
        "sudo chown syslog:adm /var/log/remote\n"
        "sudo chmod 750 /var/log/remote\n"
        "sudo systemctl restart rsyslog\n"
        "sudo ss -tulnp | grep 514"
    )
    add_body(
        doc,
        "Nếu Ubuntu bật UFW, cần mở firewall cho port 514: sudo ufw allow 514/udp và sudo ufw allow 514/tcp."
    )

    add_heading(doc, "6. Cấu hình nguồn gửi Syslog", 1)
    add_heading(doc, "6.1 Linux Server gửi Syslog", 2)
    add_code(
        doc,
        "sudo nano /etc/rsyslog.d/90-send-to-zabbix.conf\n"
        "*.* @<IP_ZABBIX_SERVER>:514\n\n"
        "sudo systemctl restart rsyslog\n"
        "logger \"TEST SYSLOG FROM LINUX SERVER\""
    )
    add_body(doc, "Ký hiệu @ gửi Syslog bằng UDP; nếu muốn dùng TCP thì dùng @@<IP_ZABBIX_SERVER>:514.")

    add_heading(doc, "6.2 Switch Cisco/IOU gửi Syslog", 2)
    add_code(
        doc,
        "enable\n"
        "conf t\n"
        "hostname SW-LAB\n"
        "service timestamps log datetime msec\n"
        "logging host <IP_ZABBIX_SERVER>\n"
        "logging trap warnings\n"
        "logging facility local7\n"
        "end\n"
        "write memory"
    )
    add_body(doc, "Tạo log test bằng cách shutdown/no shutdown một interface hoặc thay đổi mô tả cổng.")

    add_heading(doc, "6.3 Router/Firewall giả lập gửi Syslog", 2)
    add_code(
        doc,
        "enable\n"
        "conf t\n"
        "hostname R1-FW\n"
        "service timestamps log datetime msec\n"
        "logging host <IP_ZABBIX_SERVER>\n"
        "logging trap warnings\n"
        "logging facility local7\n"
        "end\n"
        "write memory"
    )
    add_body(
        doc,
        "Với firewall thật, cần cấu hình policy gửi syslog về IP Zabbix/SIEM và chọn mức log phù hợp như warning, error hoặc security event."
    )

    add_heading(doc, "7. Đưa Syslog vào Zabbix", 1)
    add_heading(doc, "7.1 Kiểm tra server đã nhận log", 2)
    add_code(doc, "sudo ls -l /var/log/remote\nsudo tail -f /var/log/remote/*\nsudo tcpdump -ni any udp port 514")

    add_heading(doc, "7.2 Cấp quyền cho Zabbix đọc log", 2)
    add_code(
        doc,
        "sudo usermod -aG adm zabbix\n"
        "sudo chmod -R 750 /var/log/remote\n"
        "sudo systemctl restart zabbix-agent\n"
        "sudo systemctl restart zabbix-server\n"
        "# Nếu dùng agent2:\n"
        "sudo systemctl restart zabbix-agent2"
    )

    add_heading(doc, "7.3 Tạo host Central Syslog", 2)
    add_table(
        doc,
        ["Trường", "Giá trị đề xuất"],
        [
            ("Host name", "Central-Syslog"),
            ("Visible name", "Central Syslog Collector"),
            ("Host group", "SNMPSim Lab hoặc Linux servers"),
            ("Interface", "Agent, IP 127.0.0.1 nếu agent chạy cùng Zabbix Server"),
            ("Monitored by", "Server"),
        ],
        widths=[Inches(2.0), Inches(4.2)],
    )

    add_heading(doc, "7.4 Tạo item đọc log", 2)
    add_table(
        doc,
        ["Nguồn log", "Item name", "Key mẫu", "Kiểu"],
        [
            ("Switch", "Switch Syslog Critical Events", 'log[/var/log/remote/SW-LAB.log,"(%LINK|DOWN|UPDOWN|CONFIG|LOGIN|FAIL|ERROR|WARNING)",,,skip]', "Log"),
            ("Router/Firewall", "Router Firewall Syslog Events", 'log[/var/log/remote/R1-FW.log,"(DOWN|CONFIG|LOGIN|FAIL|DENY|ERROR|WARNING)",,,skip]', "Log"),
            ("Linux Server", "Linux Security Syslog Events", 'log[/var/log/remote/linux-server.log,"(ERROR|WARNING|FAIL|FAILED|sudo|authentication|denied)",,,skip]', "Log"),
        ],
        widths=[Inches(1.15), Inches(1.55), Inches(3.25), Inches(0.75)],
    )
    add_callout(
        doc,
        "Khuyến nghị:",
        "Trong môi trường log rotate, nên dùng logrt[] thay cho log[] nếu tên file có thể xoay vòng hoặc thay đổi.",
        fill="F4F6F9",
    )

    add_heading(doc, "8. Bộ lọc, trigger và phân loại cảnh báo", 1)
    add_table(
        doc,
        ["Sự kiện", "Regex/điều kiện phát hiện", "Mức cảnh báo", "Ý nghĩa"],
        [
            ("Interface down", "DOWN|down|changed state to down", "Average/High", "Đường kết nối mạng bị ngắt"),
            ("Interface up", "UP|changed state to up", "Information", "Sự cố đã phục hồi hoặc cổng bật lại"),
            ("Thay đổi cấu hình", "CONFIG|configured|SYS-5-CONFIG", "Warning", "Có thay đổi cấu hình trên thiết bị"),
            ("Đăng nhập thất bại", "FAIL|FAILED|authentication failure|Invalid|denied", "High", "Có dấu hiệu dò mật khẩu hoặc truy cập trái phép"),
            ("Firewall deny/threat", "DENY|blocked|threat|attack|malware", "High/Disaster", "Sự kiện bảo mật cần xử lý gấp"),
            ("Lỗi hệ thống", "ERROR|CRITICAL|panic|failed service", "Warning/High", "Dịch vụ hoặc thiết bị phát sinh lỗi"),
        ],
        widths=[Inches(1.35), Inches(2.1), Inches(1.1), Inches(1.65)],
    )
    add_heading(doc, "8.1 Ví dụ trigger trong Zabbix", 2)
    add_code(
        doc,
        "Interface down trên Switch:\n"
        "find(/Central-Syslog/log[/var/log/remote/SW-LAB.log,\"(%LINK|DOWN|UPDOWN|CONFIG|LOGIN|FAIL|ERROR|WARNING)\",,,skip],5m,\"regexp\",\"DOWN|down|changed state to down\")=1\n\n"
        "Login thất bại trên Linux Server:\n"
        "find(/Central-Syslog/log[/var/log/remote/linux-server.log,\"(ERROR|WARNING|FAIL|FAILED|sudo|authentication|denied)\",,,skip],5m,\"regexp\",\"FAIL|FAILED|authentication failure|Invalid|denied\")=1"
    )

    add_heading(doc, "9. Cấu hình cảnh báo Web, Email, SMS, Call và Audio", 1)
    add_body(
        doc,
        "Zabbix tạo cảnh báo theo chuỗi: Trigger chuyển sang Problem, Action kiểm tra điều kiện, sau đó gửi message qua media type đã gắn cho user. "
        "Để tránh nhầm lẫn, cần hiểu rằng Zabbix không tự gọi điện hoặc gửi SMS nếu không có gateway bên ngoài."
    )
    add_table(
        doc,
        ["Kênh", "Cách triển khai trong lab", "Ghi chú chính xác"],
        [
            ("Web", "Monitoring -> Problems, Dashboard widget", "Có sẵn trong Zabbix"),
            ("Email", "Media type Email qua SMTP Gmail/App Password", "Đã phù hợp để demo chính thức"),
            ("SMS", "Webhook đến SMS Gateway hoặc Telegram Bot thay thế SMS", "Cần dịch vụ ngoài nếu muốn SMS thật"),
            ("Call", "Webhook đến Twilio Voice, PagerDuty hoặc Opsgenie", "Zabbix chỉ kích hoạt API, dịch vụ ngoài thực hiện cuộc gọi"),
            ("Audio", "Frontend notifications/sound trong User settings", "Chỉ kêu khi trình duyệt Zabbix đang mở"),
        ],
        widths=[Inches(0.9), Inches(2.75), Inches(2.55)],
    )
    add_heading(doc, "9.1 Kiểm tra lỗi gửi cảnh báo", 2)
    for item in [
        "Monitoring -> Problems: xác nhận trigger đã tạo Problem.",
        "Reports -> Action log: xem Action đã chạy chưa và lỗi cụ thể nếu gửi thất bại.",
        "Users -> Users -> Admin -> Media: user phải có media Email/Telegram/Twilio và tick đúng severity.",
        "Alerts -> Actions -> Trigger actions: action phải Enabled và Operation gửi tới đúng user/media type.",
        "Nếu gặp lỗi No media defined for user, cần thêm media vào user nhận cảnh báo và bấm Update để lưu.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "10. Giám sát SNMP bổ sung cho Router, Switch, Firewall", 1)
    add_body(
        doc,
        "Ngoài Syslog, bài demo có thể bổ sung giám sát SNMP để hiển thị chỉ số định lượng. "
        "Trong lab hiện tại, SNMPSim Dashboard giả lập Firewall, Router, Switch và Ubuntu Server. "
        "Do dịch vụ snmpd của Ubuntu thường chiếm port 161, nên nên chạy SNMPSim ở port 1161."
    )
    add_code(
        doc,
        "export SNMPSIM_ENDPOINT=\"0.0.0.0:1161\"\n"
        "python app.py\n\n"
        "snmpwalk -v2c -c fortinet udp:127.0.0.1:1161 1.3.6.1.2.1.1.5.0\n"
        "snmpwalk -v2c -c router   udp:127.0.0.1:1161 1.3.6.1.2.1.1.5.0\n"
        "snmpwalk -v2c -c switch   udp:127.0.0.1:1161 1.3.6.1.2.1.1.5.0"
    )
    add_table(
        doc,
        ["Host Zabbix", "IP/Port", "Community", "OID tiêu biểu"],
        [
            ("FortiGate-FW", "127.0.0.1:1161", "fortinet", "CPU 1.3.6.1.4.1.12356.101.4.1.3.0; RAM .4.0; Disk .6.0; Sessions .8.0"),
            ("Router-Sim", "127.0.0.1:1161", "router", "CPU 1.3.6.1.4.1.9.9.109.1.1.1.1.5.1; Load 1.3.6.1.4.1.2021.10.1.3.1"),
            ("Switch-Sim", "127.0.0.1:1161", "switch", "CPU Cisco OID; interface counters 1.3.6.1.2.1.31.1.1.1.6.<ifIndex>"),
            ("Ubuntu-Sim", "127.0.0.1:1161", "ubuntu", "UCD CPU 1.3.6.1.4.1.2021.11.9.0; RAM total/available"),
        ],
        widths=[Inches(1.2), Inches(1.3), Inches(1.0), Inches(2.7)],
    )

    add_heading(doc, "11. Kịch bản demo", 1)
    add_table(
        doc,
        ["Bước", "Thao tác", "Kết quả cần chụp/ghi nhận"],
        [
            ("1", "Kiểm tra kết nối giữa Zabbix Server và các nguồn log.", "Ping thành công, tcpdump thấy gói UDP/514 nếu có log gửi tới."),
            ("2", "Gửi log test từ Linux bằng logger.", "File /var/log/remote/<hostname>.log được cập nhật."),
            ("3", "Shutdown/no shutdown interface trên Switch hoặc Router.", "Zabbix item log nhận dòng DOWN/UPDOWN."),
            ("4", "Quan sát Monitoring -> Problems.", "Problem xuất hiện theo trigger đã cấu hình."),
            ("5", "Kiểm tra Gmail/Telegram/Action log.", "Email/Telegram gửi thành công hoặc Action log ghi trạng thái gửi."),
            ("6", "Kéo CPU/RAM trên SNMPSim Dashboard nếu demo SNMP.", "Latest data SNMP trong Zabbix thay đổi và trigger metric phát sinh."),
        ],
        widths=[Inches(0.65), Inches(3.0), Inches(2.55)],
    )

    add_heading(doc, "12. Đánh giá kết quả", 1)
    add_table(
        doc,
        ["Yêu cầu", "Kết quả mong muốn", "Trạng thái"],
        [
            ("Nhận Syslog tập trung", "Zabbix Server nhận log từ Server/Switch/Router/Firewall qua rsyslog", "Đạt khi /var/log/remote có file log"),
            ("Lọc sự kiện quan trọng", "Regex phát hiện DOWN, CONFIG, LOGIN FAIL, ERROR, DENY", "Đạt khi Latest data có log khớp rule"),
            ("Sinh cảnh báo", "Trigger tạo Problem đúng severity", "Đạt khi Monitoring -> Problems hiển thị"),
            ("Gửi cảnh báo", "Email/Telegram/SMS gateway/Call gateway nhận thông báo", "Đạt khi Reports -> Action log thành công"),
            ("Giám sát metric bổ sung", "Zabbix poll SNMP CPU/RAM/traffic từ SNMPSim", "Đạt khi Latest data có SNMP values"),
        ],
        widths=[Inches(1.55), Inches(3.25), Inches(1.4)],
    )

    add_heading(doc, "13. Nhận xét, hạn chế và hướng phát triển", 1)
    add_heading(doc, "13.1 Nhận xét", 2)
    add_body(
        doc,
        "Hệ thống đã mô phỏng đúng quy trình giám sát tập trung: nguồn log gửi Syslog về Zabbix Server, rsyslog lưu log, "
        "Zabbix đọc log, trigger phân loại sự kiện và action gửi cảnh báo. Việc bổ sung SNMP giúp dashboard có thêm chỉ số định lượng để demo trực quan."
    )
    add_heading(doc, "13.2 Hạn chế", 2)
    for item in [
        "SMS và call thật cần dịch vụ bên ngoài như SMS Gateway, Twilio, PagerDuty hoặc Opsgenie.",
        "Regex trong lab còn đơn giản, cần tinh chỉnh khi triển khai trong mạng thật để giảm false positive.",
        "Nếu số lượng thiết bị lớn, cần thiết kế lưu trữ log lâu dài bằng SIEM như Wazuh, Graylog hoặc ELK.",
        "SNMPSim chỉ là dữ liệu giả lập, phù hợp demo và kiểm thử chứ không thay thế thiết bị thật.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "13.3 Hướng phát triển", 2)
    for item in [
        "Tích hợp Zabbix với Wazuh/ELK/Graylog để tìm kiếm và tương quan log tốt hơn.",
        "Tạo webhook chuẩn cho Telegram, SMS Gateway và Twilio Voice Call.",
        "Tách rule theo từng loại thiết bị: Linux, Cisco Switch, Router, Firewall.",
        "Bổ sung dashboard tổng hợp gồm Syslog events, SNMP metrics và trạng thái gửi alert.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "14. Kết luận", 1)
    add_body(
        doc,
        "Bài lab đã xây dựng được mô hình thu thập Syslog tập trung và cảnh báo an toàn mạng bằng Zabbix. "
        "Hệ thống không chỉ nhận log mà còn áp dụng bộ lọc để phát hiện sự kiện quan trọng, hiển thị Problem trên dashboard "
        "và gửi cảnh báo qua các kênh phù hợp. Việc kết hợp thêm SNMP giúp mô hình demo đầy đủ hơn vì vừa có log sự kiện, "
        "vừa có chỉ số hiệu năng thiết bị."
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(OUT)
